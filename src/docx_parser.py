"""DOCX text extraction, mirroring pdf_parser's interface (source_name,
text, ok, error) so callers can treat PDF and DOCX resumes identically.
Never lets one malformed file crash the batch."""
import io
from dataclasses import dataclass
from pathlib import Path
from typing import Union

import docx


@dataclass
class ParsedDocX:
    source_name: str
    text: str
    ok: bool
    error: str = ""


def extract_text_from_docx(source: Union[str, Path, bytes], source_name: str = "") -> ParsedDocX:
    name = source_name or (str(source) if isinstance(source, (str, Path)) else "uploaded.docx")
    try:
        # python-docx's Document() needs a path or a file-like stream — raw
        # bytes (e.g. from Streamlit's uploader) must be wrapped first.
        docx_source = io.BytesIO(source) if isinstance(source, bytes) else source
        document = docx.Document(docx_source)
    except Exception as exc:  # noqa: BLE001 - genuinely want to catch anything python-docx raises
        return ParsedDocX(source_name=name, text="", ok=False, error=f"Could not open DOCX: {exc}")

    try:
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        full_text = "\n".join(parts)
    except Exception as exc:  # noqa: BLE001
        return ParsedDocX(source_name=name, text="", ok=False, error=f"Could not read document body: {exc}")

    if not full_text.strip():
        return ParsedDocX(source_name=name, text="", ok=False, error="No extractable text found in DOCX")

    return ParsedDocX(source_name=name, text=full_text, ok=True)
