from src.docx_parser import extract_text_from_docx
from src.document_loader import load_document
from src.pdf_parser import extract_text_from_pdf
from src.resume_parser import parse_resume_sections
from src.jd_parser import parse_jd
from src.text_cleaner import clean_text
from tests.conftest import SAMPLE_JD_TEXT, STRONG_RESUME_TEXT


def test_pdf_extraction_valid(make_pdf):
    data = make_pdf(STRONG_RESUME_TEXT)
    result = extract_text_from_pdf(data, source_name="test.pdf")
    assert result.ok
    assert "React" in result.text or "React" in result.text.replace("\n", " ")


def test_pdf_extraction_invalid_bytes_does_not_crash():
    result = extract_text_from_pdf(b"not a real pdf", source_name="broken.pdf")
    assert result.ok is False
    assert result.error


def test_docx_extraction_from_raw_bytes(tmp_path):
    """Regression test: python-docx's Document() needs a path or file-like
    stream, not raw bytes — this must be wrapped internally (previously
    caused every uploaded .docx to silently fail via the Streamlit path,
    which passes raw bytes, while a path-based script call still worked)."""
    import docx as docx_lib

    doc = docx_lib.Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Skills: React, Node.js, MongoDB")
    docx_path = tmp_path / "resume.docx"
    doc.save(str(docx_path))
    raw_bytes = docx_path.read_bytes()

    result = extract_text_from_docx(raw_bytes, source_name="resume.docx")
    assert result.ok, result.error
    assert "React" in result.text

    via_loader = load_document(raw_bytes, source_name="resume.docx")
    assert via_loader.ok
    assert "MongoDB" in via_loader.text


def test_docx_extraction_invalid_bytes_does_not_crash():
    result = extract_text_from_docx(b"not a real docx", source_name="broken.docx")
    assert result.ok is False
    assert result.error


def test_pdf_extraction_empty_pdf_flagged(make_pdf):
    data = make_pdf("")
    result = extract_text_from_pdf(data, source_name="empty.pdf")
    assert result.ok is False


def test_resume_section_detection_maps_aliases():
    text = "Jane Doe\n\nTechnical Expertise\nPython, SQL\n\nProfessional Experience\nDid things."
    parsed = parse_resume_sections(text)
    assert "skills" in parsed.sections
    assert "experience" in parsed.sections
    assert "python" in parsed.sections["skills"].lower()


def test_resume_extracts_email_and_name():
    parsed = parse_resume_sections(STRONG_RESUME_TEXT)
    assert parsed.email == "test.strong@example.com"
    assert parsed.name == "Test Candidate Strong"


def test_resume_no_headers_falls_back_to_summary():
    parsed = parse_resume_sections("Just some free text with no headers at all here.")
    assert "summary" in parsed.sections


def test_jd_parsing_extracts_required_and_preferred():
    jd = parse_jd(SAMPLE_JD_TEXT)
    assert "react" in jd.required_skills
    assert "node.js" in jd.required_skills
    assert "typescript" in jd.preferred_skills
    assert "typescript" not in jd.required_skills


def test_jd_parsing_without_headings_falls_back_to_full_text_scan():
    text = "We need someone who knows Python and Django. Nice to work with SQL too."
    jd = parse_jd(text)
    assert "python" in jd.required_skills
    assert "django" in jd.required_skills
